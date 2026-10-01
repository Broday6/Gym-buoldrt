import json
import tempfile
import unittest
from pathlib import Path

from ptk import runsheet
from ptk.doi import Lease, build_deck, participation
from ptk.ledger import load_events, replay

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

EXAMPLES = Path(__file__).resolve().parents[1] / "examples"


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class SpreadsheetExportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_runsheet_links_colors_and_flags(self):
        from ptk.xlsx import FILLS, export_runsheet

        images = self.tmp / "images"
        images.mkdir()
        rows, flags = runsheet.check_file(EXAMPLES / "runsheet.csv", images)
        out = export_runsheet(rows, flags, self.tmp / "out" / "runsheet.xlsx", images)
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Runsheet", "Flags", "About"])
        ws = wb["Runsheet"]
        headers = [c.value for c in ws[1]]
        img = ws.cell(row=2, column=headers.index("image_file") + 1)
        self.assertEqual(img.hyperlink.target, "../images/TX-REEVES-DR-0088-0077.pdf")
        self.assertEqual(ws.cell(row=2, column=1).fill.fgColor.rgb[-6:], FILLS["stop"])  # image missing -> stop
        self.assertIn("DOUBLE_FRACTION", ws.cell(row=2, column=len(headers)).value)
        self.assertEqual(wb["Flags"].max_row, len(flags) + 1)
        self.assertEqual(ws.freeze_panes, "A2")

    def test_deck_has_live_formulas(self):
        from ptk.xlsx import export_deck

        p = json.loads((EXAMPLES / "project.json").read_text())
        r = replay(load_events(EXAMPLES / "takeoff.csv"))
        weights = p["participation"]["weights"]
        deck = build_deck(r, [Lease.from_dict(d) for d in p["leases"]], participation(weights),
                          npri_ratified=p["npri_ratified"])
        out = export_deck(deck, weights, self.tmp / "deck.xlsx", title=p["well"])
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb.sheetnames, ["By owner", "Participation", "By tract", "Flags", "About"])
        self.assertTrue(wb.calculation.fullCalcOnLoad)
        self.assertEqual(wb["Participation"]["C2"].value, "=B2/SUM($B$2:$B$3)")
        self.assertEqual(wb["Participation"]["B2"].value, 4800)
        tract = wb["By tract"]
        self.assertTrue(tract["G2"].value.startswith("=INDEX(Participation!"))
        self.assertEqual(tract["H2"].value, "=E2*G2")
        owner = wb["By owner"]
        self.assertTrue(owner["C2"].value.startswith("=SUMIFS("))
        labels = [owner.cell(row=i, column=1).value for i in range(1, owner.max_row + 1)]
        self.assertIn("Check (must read OK)", labels)


if __name__ == "__main__":
    unittest.main()
