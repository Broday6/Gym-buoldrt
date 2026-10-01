"""Spreadsheet exports: the runsheet with clickable image links, and the unit deck.

Requires openpyxl (``pip install openpyxl``).

The deck workbook keeps the participation weights as live inputs. Change a
tract's acres or lateral feet on the Participation sheet and every unit decimal,
owner total and the whole-deck check recalculate. Tract decimals come from the
title work (exact fractions, shown alongside) and are values, not formulas.
"""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path

from .doi import REVENUE_TYPES, Deck
from .fracs import fmt
from .runsheet import COLUMNS, RowFlag

FONT = "Arial"
DECIMAL_FORMAT = "0.00000000"
FILLS = {"stop": "F4CCCC", "review": "FCE8B2", "info": "D9EAD3", "header": "D9DEE4", "input": "FFF2CC"}
SEVERITY_RANK = {"stop": 3, "review": 2, "info": 1}


def _openpyxl():
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:  # pragma: no cover - depends on environment
        raise SystemExit("Spreadsheet export needs openpyxl: pip install openpyxl") from e
    return openpyxl, Alignment, Font, PatternFill, get_column_letter


def _header(ws, headers: list[str], widths: dict[str, int] | None = None) -> None:
    _, Alignment, Font, PatternFill, get_column_letter = _openpyxl()
    ws.append(headers)
    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = PatternFill("solid", fgColor=FILLS["header"])
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = (widths or {}).get(name, max(12, min(len(name) + 4, 40)))
    ws.freeze_panes = "A2"


def _body_font(ws) -> None:
    _, Alignment, Font, _, _ = _openpyxl()
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.font is None or not cell.font.bold:
                cell.font = Font(name=FONT, color=cell.font.color if cell.font else None,
                                 underline=cell.font.underline if cell.font else None)
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _about(wb, lines: list[str]) -> None:
    _, _, Font, _, _ = _openpyxl()
    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line).font = Font(name=FONT, bold=(i == 1))


def export_runsheet(rows: list[dict[str, str]], flags: list[RowFlag], out: str | Path,
                    images_dir: str | Path | None = None) -> Path:
    """Write the runsheet, a flags sheet, and a legend. Image file names become links."""
    openpyxl, _, Font, PatternFill, get_column_letter = _openpyxl()
    out = Path(out)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Runsheet"
    headers = COLUMNS + ["flags"]
    widths = {"grantor": 28, "grantee": 28, "legal_description": 44, "interest_conveyed": 40,
              "reservations_exceptions": 40, "notes": 32, "flags": 34, "image_file": 30}
    _header(ws, headers, widths)

    worst: dict[str, str] = {}
    codes: dict[str, list[str]] = {}
    for f in flags:
        codes.setdefault(f.row, []).append(f.code)
        if SEVERITY_RANK[f.severity] > SEVERITY_RANK.get(worst.get(f.row, ""), 0):
            worst[f.row] = f.severity

    image_col = headers.index("image_file") + 1
    for r_idx, row in enumerate(rows, start=2):
        rid = row.get("row") or ""
        ws.append([row.get(c, "") for c in COLUMNS] + [", ".join(codes.get(rid, []))])
        image = (row.get("image_file") or "").strip()
        if image and images_dir is not None:
            target = Path(images_dir) / image
            cell = ws.cell(row=r_idx, column=image_col)
            cell.hyperlink = os.path.relpath(target, out.parent).replace(os.sep, "/")
            cell.font = Font(name=FONT, color="0563C1", underline="single")
        sev = worst.get(rid)
        if sev:
            fill = PatternFill("solid", fgColor=FILLS[sev])
            for c in range(1, len(headers) + 1):
                ws.cell(row=r_idx, column=c).fill = fill
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    _body_font(ws)

    fs = wb.create_sheet("Flags")
    _header(fs, ["row", "severity", "code", "message"], {"message": 100, "code": 26})
    for f in sorted(flags, key=lambda x: (-SEVERITY_RANK[x.severity], x.row)):
        fs.append([f.row, f.severity, f.code, f.message])
        fs.cell(row=fs.max_row, column=2).fill = PatternFill("solid", fgColor=FILLS[f.severity])
    fs.auto_filter.ref = f"A1:D{max(len(flags) + 1, 1)}"
    _body_font(fs)

    counts = {s: sum(1 for f in flags if f.severity == s) for s in ("stop", "review", "info")}
    _about(wb, [
        "Runsheet export (landman work product, not a title opinion)",
        f"Generated {date.today().isoformat()} by the Permian title kit.",
        f"{len(rows)} instruments; {counts['stop']} stop, {counts['review']} review and {counts['info']} info flags.",
        "Row colors: red = at least one stop flag (don't rely on the row until it's fixed);",
        "amber = review flag (a person must look); green = information only.",
        "Image file names are links to the recorded image, relative to this workbook's folder.",
        "The Flags sheet lists every flag with the reason; filter it by severity or code.",
    ])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def _save_with_recalc(wb, out: Path) -> None:
    from openpyxl.workbook.properties import CalcProperties

    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # openpyxl stores no results; Excel computes on open
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def export_deck(deck: Deck, weights: dict[str, object], out: str | Path, title: str = "Unit deck") -> Path:
    """Write the deck with live participation formulas, owner totals and a whole-deck check."""
    openpyxl, _, Font, PatternFill, _ = _openpyxl()
    from .fracs import parse_fraction

    out = Path(out)
    wb = openpyxl.Workbook()

    # Participation: weights are the inputs; factors are formulas.
    ps = wb.active
    ps.title = "Participation"
    _header(ps, ["tract", "weight (acres in unit or completed lateral feet)", "factor"],
            {"weight (acres in unit or completed lateral feet)": 30, "factor": 16})
    tracts = list(weights)
    n = len(tracts)
    for i, tract in enumerate(tracts, start=2):
        ps.cell(row=i, column=1, value=tract)
        w = ps.cell(row=i, column=2, value=float(parse_fraction(str(weights[tract]))))
        w.font = Font(name=FONT, color="0000FF")
        w.fill = PatternFill("solid", fgColor=FILLS["input"])
        f = ps.cell(row=i, column=3, value=f"=B{i}/SUM($B$2:$B${n + 1})")
        f.number_format = DECIMAL_FORMAT
    ps.cell(row=n + 2, column=1, value="Total")
    ps.cell(row=n + 2, column=2, value=f"=SUM(B2:B{n + 1})")
    ps.cell(row=n + 2, column=3, value=f"=SUM(C2:C{n + 1})").number_format = DECIMAL_FORMAT
    _body_font(ps)
    for c in (1, 2, 3):
        ps.cell(row=n + 2, column=c).font = Font(name=FONT, bold=True)

    # By tract: tract decimals are values from the title work; unit decimals are formulas.
    ts = wb.create_sheet("By tract")
    _header(ts, ["owner", "type", "tract", "depth", "tract decimal", "tract exact", "factor", "unit decimal",
                 "source", "note"], {"owner": 30, "source": 26, "note": 52})
    for i, r in enumerate(deck.rows, start=2):
        ts.append([r.owner, r.type, r.tract, r.depth, float(r.tract_decimal), fmt(r.tract_decimal), None, None,
                   r.source, r.note])
        ts.cell(row=i, column=7, value=f"=INDEX(Participation!$C$2:$C${n + 1},MATCH(C{i},Participation!$A$2:$A${n + 1},0))")
        ts.cell(row=i, column=8, value=f"=E{i}*G{i}")
        for c in (5, 7, 8):
            ts.cell(row=i, column=c).number_format = DECIMAL_FORMAT
    last = len(deck.rows) + 1
    ts.auto_filter.ref = f"A1:J{max(last, 1)}"
    _body_font(ts)

    # By owner: SUMIFS over the tract rows.
    os_ = wb.create_sheet("By owner", 0)
    _header(os_, ["owner", "type", "unit decimal", "exact (from title work)"], {"owner": 32, "exact (from title work)": 24})
    owners = deck.by_owner()
    for i, (owner, typ, value) in enumerate(owners, start=2):
        os_.cell(row=i, column=1, value=owner)
        os_.cell(row=i, column=2, value=typ)
        c = os_.cell(row=i, column=3, value=f"=SUMIFS('By tract'!$H$2:$H${last},'By tract'!$A$2:$A${last},A{i},"
                                           f"'By tract'!$B$2:$B${last},B{i})")
        c.number_format = DECIMAL_FORMAT
        os_.cell(row=i, column=4, value=fmt(value))
    end = len(owners) + 1
    row = end + 2
    os_.cell(row=row, column=1, value="Revenue interests total")
    terms = "+".join(f'SUMIFS(C2:C{end},B2:B{end},"{t}")' for t in REVENUE_TYPES)
    total = os_.cell(row=row, column=3, value=f"={terms}")
    total.number_format = DECIMAL_FORMAT
    os_.cell(row=row + 1, column=1, value="Check (must read OK)")
    os_.cell(row=row + 1, column=3, value=f'=IF(ROUND(C{row},8)=1,"OK","NOT WHOLE")')
    _body_font(os_)
    for c in (1, 3):
        os_.cell(row=row, column=c).font = Font(name=FONT, bold=True)
        os_.cell(row=row + 1, column=c).font = Font(name=FONT, bold=True)
    os_.freeze_panes = "A2"

    fl = wb.create_sheet("Flags")
    _header(fl, ["code", "tract", "depth", "instrument", "detail"], {"detail": 100, "instrument": 30, "code": 26})
    for f in deck.flags:
        fl.append([f.code, f.tract, f.depth, f.instrument, f.detail])
    _body_font(fl)

    exact_total = deck.revenue_total()
    _about(wb, [
        f"{title} (landman work product, not a title opinion)",
        f"Generated {date.today().isoformat()} by the Permian title kit.",
        "Blue cells on a yellow fill (Participation, column B) are inputs: acres in the unit or completed lateral feet per tract.",
        "Everything else on Participation, By tract and By owner recalculates from them.",
        "Tract decimals are exact results from the title work; the 'exact' columns show the fraction.",
        f"Revenue total from the title work: {fmt(exact_total)}. The By owner check reads OK only when the deck adds to 1.",
        f"{len(deck.flags)} flags; see the Flags sheet before relying on any decimal.",
    ])
    _save_with_recalc(wb, out)
    return out

